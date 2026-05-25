"""DOCX paragraph parsing helpers for structure checks."""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.oxml.ns import qn


@dataclass
class DocPara:
    text: str
    style: str
    heading_level: int


def _norm_ws(s: str) -> str:
    """Normalize whitespace while preserving paragraph boundaries."""
    s = (s or "").replace("\u00A0", " ")
    s = re.sub(r"[ \t]+", " ", s).strip()
    s = re.sub(r"\n{3,}", "\n\n", s).strip()
    return s


def _collect_header_footer_text(doc: Document) -> set[str]:
    hf_texts: set[str] = set()
    try:
        for sec in doc.sections:
            for p in sec.header.paragraphs:
                t = _norm_ws(p.text)
                if t:
                    hf_texts.add(t)
            for p in sec.footer.paragraphs:
                t = _norm_ws(p.text)
                if t:
                    hf_texts.add(t)
    except Exception:
        pass
    return hf_texts


def _collect_numbering_formats(doc: Document) -> Dict[Tuple[str, str], Tuple[str, str, int]]:
    """Map (numId, ilvl) -> (numFmt, lvlText, start)."""
    try:
        root = doc.part.numbering_part.element
    except Exception:
        return {}

    num_to_abs: Dict[str, str] = {}
    for num in root.findall(qn("w:num")):
        num_id = num.get(qn("w:numId"))
        abs_el = num.find(qn("w:abstractNumId"))
        if num_id is None or abs_el is None:
            continue
        abs_id = abs_el.get(qn("w:val"))
        if abs_id is not None:
            num_to_abs[num_id] = abs_id

    abs_formats: Dict[Tuple[str, str], Tuple[str, str, int]] = {}
    for abs_num in root.findall(qn("w:abstractNum")):
        abs_id = abs_num.get(qn("w:abstractNumId"))
        if abs_id is None:
            continue
        for lvl in abs_num.findall(qn("w:lvl")):
            ilvl = lvl.get(qn("w:ilvl"))
            fmt_el = lvl.find(qn("w:numFmt"))
            text_el = lvl.find(qn("w:lvlText"))
            start_el = lvl.find(qn("w:start"))
            if ilvl is None or fmt_el is None or text_el is None:
                continue
            fmt = fmt_el.get(qn("w:val")) or ""
            text = text_el.get(qn("w:val")) or ""
            try:
                start = int(start_el.get(qn("w:val"))) if start_el is not None else 1
            except Exception:
                start = 1
            abs_formats[(abs_id, ilvl)] = (fmt, text, start)

    out: Dict[Tuple[str, str], Tuple[str, str, int]] = {}
    for num_id, abs_id in num_to_abs.items():
        for (fmt_abs_id, ilvl), spec in abs_formats.items():
            if fmt_abs_id == abs_id:
                out[(num_id, ilvl)] = spec
    return out


def _paragraph_numbering_key(p) -> Optional[Tuple[str, str]]:
    ppr = p._p.pPr
    if ppr is None or ppr.numPr is None:
        return None
    num_pr = ppr.numPr
    if num_pr.numId is None or num_pr.ilvl is None:
        return None
    return str(num_pr.numId.val), str(num_pr.ilvl.val)


def _format_numbering_label(lvl_text: str, number: int) -> str:
    return (lvl_text or "%1").replace("%1", str(number))


_HEADING_STYLE_RE = re.compile(r"^(heading|tiêu đề|tieu de)\s+(\d+)\s*$", re.I)


def _heading_level_from_style(style_name: str) -> int:
    s = (style_name or "").strip()
    m = _HEADING_STYLE_RE.match(s)
    if not m:
        return 0
    try:
        return int(m.group(2))
    except Exception:
        return 0


def extract_docx_paragraphs(path: str) -> List[DocPara]:
    """Read DOCX into normalized non-empty body paragraphs."""
    doc = Document(path)
    hf_texts = _collect_header_footer_text(doc)
    numbering_formats = _collect_numbering_formats(doc)
    numbering_counters: Dict[Tuple[str, str], int] = defaultdict(int)

    out: List[DocPara] = []
    for p in doc.paragraphs:
        t = _norm_ws(p.text)
        if not t:
            continue
        if t in hf_texts:
            continue

        try:
            style_name = p.style.name or ""
        except Exception:
            style_name = ""

        lvl = _heading_level_from_style(style_name)
        numbering_key = _paragraph_numbering_key(p)
        if numbering_key and numbering_key in numbering_formats:
            num_fmt, lvl_text, start = numbering_formats[numbering_key]
            if num_fmt == "decimal" and "%1" in lvl_text:
                if numbering_counters[numbering_key] == 0:
                    numbering_counters[numbering_key] = start
                else:
                    numbering_counters[numbering_key] += 1
                label = _format_numbering_label(lvl_text, numbering_counters[numbering_key])
                if label and not t.startswith(label):
                    t = f"{label} {t}"

        out.append(DocPara(text=t, style=style_name, heading_level=lvl))

    return out


def paragraphs_to_full_text(paras: List[DocPara]) -> str:
    """Join parsed paragraphs into the full paper text sent to the LLM."""
    chunks = [_norm_ws(p.text) for p in paras if p.text and p.text.strip()]
    return _norm_ws("\n".join(chunk for chunk in chunks if chunk))
